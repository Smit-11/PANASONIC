import pyodbc
import time
import threading
import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
from tkcalendar import DateEntry

# ---------------- DATABASE CONFIG ----------------
server = r'Admin\SQLEXPRESS'
database = 'SI850'
uid = 'sa'
pwd = 'password123'

connection_string = (
    f'DRIVER={{ODBC Driver 18 for SQL Server}};'
    f'SERVER={server};'
    f'DATABASE={database};'
    f'UID={uid};'
    f'PWD={pwd};'
    'Encrypt=no;'
    'TrustServerCertificate=yes;'
)

# ---------------- SETTINGS ----------------
OFFLINE_DELAY = 120
LIVE_REFRESH_MS = 3000
LIVE_RECORD_LIMIT = 1000

# ---------------- GLOBAL STORAGE ----------------
loaded_records = set()
machine_status = {}
machine_timer = {}
machine_last_remark = {}
live_records = []
machine_vars = {}
machine_cbs = {}
download_machine_vars = {}
download_machine_cbs = {}
search_btn = None
clear_btn = None
live_status_label = None
MACHINE_NAMES = ["C-1","C-3","C-4","C-5","C-6","C-7","C-8","SM","R1","R2","R1R2"]

table_meta_cache = {}
live_connection = None
latest_live_ids = {}
live_fetch_lock = threading.Lock()
live_fetching = False

def normalize_machine(name: str) -> str:
    if not name:
        return ""
    # Remove non-alphanumeric characters and uppercase
    try:
        s = str(name)
    except:
        s = name or ""
    return re.sub(r'[^A-Za-z0-9]', '', s).upper()


def quote_identifier(name: str) -> str:
    safe_name = str(name).replace(']', ']]')
    return f'[{safe_name}]'


def safe_config(widget, **kwargs):
    if widget is not None:
        try:
            widget.config(**kwargs)
        except Exception:
            pass


def parse_datetime_field(value):
    if isinstance(value, datetime):
        return value
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None

    formats = [
        '%d-%m-%Y %H:%M:%S',
        '%d-%m-%Y %H:%M',
        '%d-%m-%Y',
        '%Y-%m-%d %H:%M:%S',
        '%Y-%m-%d %H:%M',
        '%Y-%m-%d',
        '%d/%m/%Y %H:%M:%S',
        '%d/%m/%Y %H:%M',
        '%d/%m/%Y',
        '%m/%d/%Y %H:%M:%S',
        '%m/%d/%Y %H:%M',
        '%m/%d/%Y',
    ]
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except Exception:
            continue
    return None


def get_row_datetime(row, columns, date_col, time_col=None):
    if not date_col:
        return None
    date_value = row[columns.index(date_col)]
    if isinstance(date_value, datetime):
        base = date_value
    else:
        base = parse_datetime_field(date_value)

    if not base:
        return None

    if time_col:
        time_value = row[columns.index(time_col)]
        if time_value is None:
            return base

        parsed_time = parse_datetime_field(time_value)
        if isinstance(parsed_time, datetime):
            return datetime.combine(base.date(), parsed_time.time())

        time_text = str(time_value).strip()
        for fmt in ('%H:%M:%S', '%H:%M', '%I:%M:%S %p', '%I:%M %p'):
            try:
                parsed = datetime.strptime(time_text, fmt)
                return datetime.combine(base.date(), parsed.time())
            except Exception:
                continue

    return base


def get_table_metadata(cursor, table):
    if table in table_meta_cache:
        return table_meta_cache[table]

    cursor.execute(f"SELECT TOP 1 * FROM {table}")
    columns = [col[0] for col in cursor.description]

    machine_col = None
    possible_machine_columns = ["M/C NO04", "MACHINE NAME04", "MNo"]
    for col in possible_machine_columns:
        if col in columns:
            machine_col = col
            break

    if not machine_col:
        return None

    remark_col = None
    for col in columns:
        if "REMARK" in col.upper():
            remark_col = col
            break

    date_col = None
    for col in columns:
        if "DATE" in col.upper() or "DATETIME" in col.upper():
            date_col = col
            break

    time_col = None
    for col in columns:
        if "TIME" in col.upper() and col != date_col:
            time_col = col
            break

    rec_no_col = None
    for col in columns:
        if col.upper().replace("_", "").replace(" ", "") == "RECNO" or col.upper().replace("_", "").replace(" ", "") == "RECNO":
            rec_no_col = col
            break
    if not rec_no_col:
        for col in columns:
            if "REC" in col.upper() and "NO" in col.upper():
                rec_no_col = col
                break
    if not rec_no_col:
        rec_no_col = columns[0]

    metadata = {
        "columns": columns,
        "machine_col": machine_col,
        "remark_col": remark_col,
        "remark_idx": columns.index(remark_col) if remark_col else None,
        "date_col": date_col,
        "time_col": time_col,
        "rec_no_col": rec_no_col,
        "rec_no_idx": columns.index(rec_no_col),
    }
    table_meta_cache[table] = metadata
    return metadata

# ---------------- FETCH NEW DATA ONLY ----------------
def fetch_new_data():
    connection = ensure_live_connection()
    if connection is None:
        return None
    cursor = connection.cursor()

    cursor.execute("""
        SELECT name
        FROM sys.tables
        WHERE name LIKE 'MAC_Rec%'
    """)

    tables = [row.name for row in cursor.fetchall()]
    all_new_data = []

    for table in tables:
        try:
            meta = get_table_metadata(cursor, table)
            if not meta:
                continue

            columns = meta["columns"]
            machine_col = meta["machine_col"]
            remark_idx = meta["remark_idx"]
            date_col = meta["date_col"]
            time_col = meta["time_col"]
            rec_no_col = meta["rec_no_col"]
            rec_no_idx = meta["rec_no_idx"]
            quoted_table = quote_identifier(table)
            quoted_rec_no = quote_identifier(rec_no_col)

            if table in latest_live_ids and str(latest_live_ids[table]).isdigit():
                query = f"""
                    SELECT TOP 50 *
                    FROM {quoted_table}
                    WHERE TRY_CONVERT(bigint, {quoted_rec_no}) > ?
                    ORDER BY TRY_CONVERT(bigint, {quoted_rec_no}) DESC
                """
                cursor.execute(query, latest_live_ids[table])
            else:
                query = f"""
                    SELECT TOP 50 *
                    FROM {quoted_table}
                    ORDER BY TRY_CONVERT(bigint, {quoted_rec_no}) DESC
                """
                cursor.execute(query)

            rows = cursor.fetchall()

            max_seen = latest_live_ids.get(table)
            for row in rows:
                rec_no = str(row[rec_no_idx])
                if not rec_no:
                    continue
                if rec_no.isdigit():
                    rec_int = int(rec_no)
                    if max_seen is None or rec_int > int(max_seen):
                        max_seen = rec_no

                key = (table, rec_no)
                if key not in loaded_records:
                    record = tuple(str(x) for x in row)
                    machine_no = str(row[columns.index(machine_col)])
                    sort_key = get_row_datetime(row, columns, date_col, time_col)
                    if sort_key is None:
                        sort_key = datetime.min.replace(year=1900) + (int(rec_no) if rec_no.isdigit() else 0)
                    all_new_data.append((table, rec_no, machine_no, record, remark_idx, columns, sort_key))

            if max_seen is not None:
                latest_live_ids[table] = max_seen

        except Exception as e:
            print(f"ERROR in table {table}: {e}")

    cursor.close()

    all_new_data.sort(key=lambda x: (x[6] is not None, x[6] or datetime.min), reverse=True)
    if len(all_new_data) > LIVE_RECORD_LIMIT:
        all_new_data = all_new_data[:LIVE_RECORD_LIMIT]
    
    return all_new_data

# ---------------- FETCH RECORDS BY DATE RANGE ----------------
def fetch_records_by_date(start_date, end_date):
    connection = pyodbc.connect(connection_string)
    cursor = connection.cursor()

    cursor.execute("""
        SELECT name
        FROM sys.tables
        WHERE name LIKE 'MAC_Rec%'
    """)

    tables = [row.name for row in cursor.fetchall()]
    all_records = []

    for table in tables:
        try:
            meta = get_table_metadata(cursor, table)
            if not meta:
                continue

            columns = meta["columns"]
            machine_col = meta["machine_col"]
            remark_idx = meta["remark_idx"]
            date_col = meta["date_col"]
            time_col = meta["time_col"]
            rec_no_col = meta["rec_no_col"]
            rec_no_idx = meta["rec_no_idx"]
            quoted_table = quote_identifier(table)
            quoted_date_col = quote_identifier(date_col) if date_col else None
            quoted_rec_no = quote_identifier(rec_no_col)

            query = f"""
                SELECT *
                FROM {quoted_table}
                WHERE 1=1
            """

            if date_col:
                query += (
                    " AND ("
                    f"TRY_CONVERT(date, {quoted_date_col}, 105) BETWEEN TRY_CONVERT(date, '{start_date}', 105) AND TRY_CONVERT(date, '{end_date}', 105)"
                    " OR "
                    f"TRY_CONVERT(date, {quoted_date_col}) BETWEEN TRY_CONVERT(date, '{start_date}', 105) AND TRY_CONVERT(date, '{end_date}', 105)"
                    ")"
                )

            query += f" ORDER BY {quoted_rec_no} DESC"
            cursor.execute(query)
            rows = cursor.fetchall()

            for row in rows:
                rec_no = str(row[rec_no_idx])
                if not rec_no:
                    continue
                record = tuple(str(x) for x in row)
                machine_no = str(row[columns.index(machine_col)])

                date_value = ""
                if date_col:
                    raw_val = row[columns.index(date_col)]
                    if isinstance(raw_val, datetime):
                        date_value = raw_val.strftime('%d-%m-%Y')
                    else:
                        raw_s = str(raw_val)
                        parsed = None
                        for fmt in ('%d-%m-%Y', '%m-%d-%Y', '%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%Y-%m-%d %H:%M:%S'):
                            try:
                                parsed = datetime.strptime(raw_s, fmt)
                                break
                            except Exception:
                                parsed = None
                        if parsed:
                            date_value = parsed.strftime('%d-%m-%Y')
                        else:
                            date_value = raw_s

                sort_key = get_row_datetime(row, columns, date_col, time_col)
                if sort_key is None and rec_no.isdigit():
                    sort_key = datetime.min.replace(year=1900) + (int(rec_no) - 1)

                all_records.append((rec_no, machine_no, record, remark_idx, columns, date_value, sort_key))

        except Exception as e:
            print(f"ERROR in table {table}: {e}")

    connection.close()

    all_records.sort(key=lambda x: (x[6] is not None, x[6] or datetime.min), reverse=True)
    return all_records

# ---------------- EXPORT TO EXCEL ----------------
def export_to_excel(records, start_date, end_date):
    if not records:
        messagebox.showwarning("No Records", "No records found in the selected date range!")
        return
    
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Machine Records"
        
        ok_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
        low_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        high_fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        if records:
            columns = records[0][4]
            header = ["Rec_NO", "Machine"] + list(columns)
            
            for col_idx, col_name in enumerate(header, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border
        
        row_num = 2
        for rec_no, machine_no, record, remark_idx, columns, date_value, sort_key in records:
            values = [rec_no, machine_no] + list(record)
            
            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')
                
                if remark_idx is not None and col_idx > 2:
                    try:
                        remark_value = str(record[remark_idx]).upper()
                        if remark_value == "OK":
                            cell.fill = ok_fill
                        elif remark_value == "LOW":
                            cell.fill = low_fill
                        elif remark_value == "HIGH":
                            cell.fill = high_fill
                    except:
                        pass
            
            row_num += 1
        
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        filename = f"Machine_Records_{start_date}_to_{end_date}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        from tkinter import filedialog
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            initialfile=filename
        )
        
        if file_path:
            wb.save(file_path)
            messagebox.showinfo("Success", f"Records exported successfully to:\n{file_path}")
        
    except Exception as e:
        messagebox.showerror("Export Error", f"Error exporting to Excel: {str(e)}")

# ---------------- LIVE TAB ----------------
def ensure_live_connection():
    global live_connection
    try:
        if live_connection is not None:
            live_connection.cursor().execute("SELECT 1")
            return live_connection
    except Exception:
        try:
            live_connection.close()
        except Exception:
            pass
        live_connection = None

    try:
        live_connection = pyodbc.connect(connection_string, autocommit=True)
        return live_connection
    except Exception as e:
        print(f"Live connection failed: {e}")
        live_connection = None
        return None


def setup_live_columns(tree, columns):
    tree["columns"] = ["Machine"] + columns
    tree["show"] = "headings"
    for col in tree["columns"]:
        width = 160 if col == "Machine" else 120
        tree.heading(col, text=col)
        tree.column(col, width=width, anchor="center", stretch=False)


def rebuild_live_tree(tree):
    if not live_records:
        return

    display_records = sorted(live_records[:LIVE_RECORD_LIMIT], key=lambda x: (x[6] is not None, x[6] or datetime.min), reverse=True)
    if not tree["columns"]:
        setup_live_columns(tree, display_records[0][5])

    tree.delete(*tree.get_children())
    selected = {normalize_machine(name) for name, var in machine_vars.items() if var.get()}
    current_time = time.time()

    for table, rec_no, machine_no, record, remark_idx, columns, sort_key in display_records:
        if selected and normalize_machine(machine_no) not in selected:
            continue

        machine_timer[machine_no] = current_time
        machine_status[machine_no] = "Online"

        tag = ""
        remark_value = ""
        if remark_idx is not None:
            remark_value = str(record[remark_idx]).upper()
            if remark_value == "OK":
                tag = "ok"
            elif remark_value == "LOW":
                tag = "low"
            elif remark_value == "HIGH":
                tag = "high"

        tree.insert("", "end", values=(machine_no,) + record, tags=(tag,))
        machine_last_remark[machine_no] = remark_value


def apply_live_new_data(tree, new_data):
    global live_records
    if not new_data:
        return

    current_time = time.time()
    new_entries = []
    for table, rec_no, machine_no, record, remark_idx, columns, sort_key in new_data:
        key = (table, rec_no)
        if key not in loaded_records:
            loaded_records.add(key)
            live_records.append((table, rec_no, machine_no, record, remark_idx, columns, sort_key))
            new_entries.append((table, rec_no, machine_no, record, remark_idx, columns, sort_key))

    if not new_entries:
        return

    live_records.sort(key=lambda x: (x[6] is not None, x[6] or datetime.min), reverse=True)
    if len(live_records) > LIVE_RECORD_LIMIT:
        live_records = live_records[:LIVE_RECORD_LIMIT]

    if not tree["columns"]:
        setup_live_columns(tree, live_records[0][5])

    selected = {normalize_machine(name) for name, var in machine_vars.items() if var.get()}
    # Insert older entries first so newer rows stay at the top when inserted at index 0.
    new_entries = sorted(new_entries, key=lambda x: (x[6] is not None, x[6] or datetime.min))
    for table, rec_no, machine_no, record, remark_idx, columns, sort_key in new_entries:
        if selected and normalize_machine(machine_no) not in selected:
            continue

        machine_timer[machine_no] = current_time
        machine_status[machine_no] = "Online"

        tag = ""
        remark_value = ""
        if remark_idx is not None:
            remark_value = str(record[remark_idx]).upper()
            if remark_value == "OK":
                tag = "ok"
            elif remark_value == "LOW":
                tag = "low"
            elif remark_value == "HIGH":
                tag = "high"

        item = tree.insert("", 0, values=(machine_no,) + record, tags=("flash",))
        if len(tree.get_children()) > LIVE_RECORD_LIMIT:
            oldest = tree.get_children()[-1]
            tree.delete(oldest)
        tree.after(500, lambda i=item, t=tag: tree.item(i, tags=(t,)))
        machine_last_remark[machine_no] = remark_value


def update_live(tree, force_rebuild=False):
    global live_fetching
    if force_rebuild:
        rebuild_live_tree(tree)
        return

    if live_fetching:
        tree.after(100, lambda: update_live(tree))
        return

    def live_background_fetch():
        global live_fetching
        with live_fetch_lock:
            live_fetching = True
            try:
                new_data = fetch_new_data()
            except Exception as e:
                print(f"Live fetch error: {e}")
                new_data = None
            finally:
                live_fetching = False

        if new_data is None:
            if live_status_label is not None:
                tree.after(0, lambda: live_status_label.config(text="Live DB connection failed", fg="red"))
        else:
            if live_status_label is not None:
                tree.after(0, lambda: live_status_label.config(text="Live feed active", fg="#1a3c6e"))
            if new_data:
                tree.after(0, lambda: apply_live_new_data(tree, new_data))

        tree.after(LIVE_REFRESH_MS, lambda: update_live(tree))

    threading.Thread(target=live_background_fetch, daemon=True).start()

# ---------------- DOWNLOAD TAB ----------------
def complete_download_records(records):
    preview_tree.delete(*preview_tree.get_children())

    if not records:
        messagebox.showinfo("No Records", "No records found in the selected date range!")
        status_label.config(text="No records found", foreground="orange")
        export_button.config(state="disabled")
    else:
        records = sorted(records, key=lambda x: (x[6] is not None, x[6] or datetime.min), reverse=True)
        columns = records[0][4]
        preview_tree["columns"] = ["Rec_NO", "Machine"] + list(columns)
        preview_tree["show"] = "headings"
        for col in preview_tree["columns"]:
            preview_tree.heading(col, text=col)
            preview_tree.column(col, width=100, anchor="center", stretch=False)

        display_records = records
        if len(records) > LIVE_RECORD_LIMIT:
            display_records = records[:LIVE_RECORD_LIMIT]

        for rec_no, machine_no, record, remark_idx, columns, date_value, sort_key in display_records:
            values = [rec_no, machine_no] + list(record)
            tag = ""
            if remark_idx is not None:
                remark_value = str(record[remark_idx]).upper()
                if remark_value == "OK":
                    tag = "ok"
                elif remark_value == "LOW":
                    tag = "low"
                elif remark_value == "HIGH":
                    tag = "high"
            preview_tree.insert("", "end", values=values, tags=(tag,))

        preview_tree.tag_configure("ok", background="#92D050")
        preview_tree.tag_configure("low", background="#FFC000")
        preview_tree.tag_configure("high", background="#FF6B6B")

        export_button.config(state="normal")
        preview_tree.records = records
        status_label.config(
        text=f"✅ Found {len(records)} records",
        foreground="green"
        )

    search_btn.config(state="normal")
    clear_btn.config(state="normal")


def download_records():
    start_dt = start_cal.get_date()
    end_dt = end_cal.get_date()

    if not start_dt or not end_dt:
        messagebox.showwarning("Invalid Input", "Please select both start and end dates!")
        return

    if start_dt > end_dt:
        messagebox.showwarning("Invalid Date Range", "Start date must be before end date!")
        return

    start_date = start_dt.strftime('%d-%m-%Y')
    end_date = end_dt.strftime('%d-%m-%Y')
    selected = {normalize_machine(name) for name, var in download_machine_vars.items() if var.get()}

    status_label.config(text="Searching records...", foreground="#1a3c6e")
    safe_config(search_btn, state="disabled")
    safe_config(clear_btn, state="disabled")
    safe_config(export_button, state="disabled")

    def download_thread():
        try:
            records = fetch_records_by_date(start_date, end_date)
            if selected:
                records = [r for r in records if normalize_machine(r[1]) in selected]
            preview_tree.after(0, lambda: complete_download_records(records))
        except Exception as e:
            preview_tree.after(0, lambda: (messagebox.showerror("Error", f"Error fetching records: {str(e)}"), status_label.config(text="❌ Error fetching records", foreground="red"), safe_config(search_btn, state="normal"), safe_config(clear_btn, state="normal")))

    threading.Thread(target=download_thread, daemon=True).start()


def export_previewed_records():
    if not hasattr(preview_tree, 'records') or not preview_tree.records:
        messagebox.showwarning("No Records", "No records to export!")
        return

    from tkinter import filedialog
    start_date = start_cal.get_date().strftime('%d-%m-%Y')
    end_date = end_cal.get_date().strftime('%d-%m-%Y')
    filename = f"Machine_Records_{start_date}_to_{end_date}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    file_path = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
        initialfile=filename
    )

    if not file_path:
        return

    status_label.config(text="Exporting records...", foreground="#1a3c6e")
    safe_config(export_button, state="disabled")
    safe_config(search_btn, state="disabled")
    safe_config(clear_btn, state="disabled")

    def export_thread():
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Machine Records"

            ok_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
            low_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
            high_fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            records = preview_tree.records
            columns = records[0][4]
            header = ["Rec_NO", "Machine"] + list(columns)

            for col_idx, col_name in enumerate(header, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border

            row_num = 2
            for rec_no, machine_no, record, remark_idx, columns, date_value, sort_key in records:
                values = [rec_no, machine_no] + list(record)
                for col_idx, value in enumerate(values, 1):
                    cell = ws.cell(row=row_num, column=col_idx, value=value)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    if remark_idx is not None and col_idx > 2:
                        try:
                            remark_value = str(record[remark_idx]).upper()
                            if remark_value == "OK":
                                cell.fill = ok_fill
                            elif remark_value == "LOW":
                                cell.fill = low_fill
                            elif remark_value == "HIGH":
                                cell.fill = high_fill
                        except Exception:
                            pass
                row_num += 1

            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            wb.save(file_path)
            preview_tree.after(0, lambda: (messagebox.showinfo("Success", f"Records exported successfully to:\n{file_path}"), status_label.config(text="Export completed", foreground="#28A745"), safe_config(search_btn, state="normal"), safe_config(clear_btn, state="normal"), safe_config(export_button, state="normal")))
        except Exception as e:
            preview_tree.after(0, lambda: (messagebox.showerror("Export Error", f"Error exporting to Excel: {str(e)}"), status_label.config(text="❌ Export failed", foreground="red"), safe_config(search_btn, state="normal"), safe_config(clear_btn, state="normal"), safe_config(export_button, state="normal")))

    threading.Thread(target=export_thread, daemon=True).start()

def clear_records():
    preview_tree.delete(*preview_tree.get_children())
    if hasattr(preview_tree, 'records'):
        del preview_tree.records
    export_button.config(state="disabled")
    status_label.config(text="Cleared", foreground="gray")
    safe_config(search_btn, state="normal")
    safe_config(clear_btn, state="normal")

# ---------------- MAIN WINDOW ----------------
def open_main_window():
    global start_cal, end_cal, preview_tree, export_button, status_label
    global search_btn, clear_btn, live_status_label
    
    root = tk.Tk()
    root.title("Machine Monitoring System")
    root.geometry("1400x700")
    root.configure(bg='#f0f0f0')
    
    style = ttk.Style()
    style.theme_use('clam')
    style.configure('TNotebook', background='#f0f0f0')
    style.configure('TNotebook.Tab', padding=[10, 5], font=('Segoe UI', 10, 'bold'))
    style.map('TNotebook.Tab', 
              background=[('selected', '#4472C4'), ('active', '#B4C6E7')],
              foreground=[('selected', 'white'), ('active', 'black')])

    notebook = ttk.Notebook(root)
    notebook.pack(fill="both", expand=True, padx=5, pady=5)

    # ---------------- LIVE TAB ----------------
    live_frame = ttk.Frame(notebook)
    notebook.add(live_frame, text="📊 Live Monitoring")

    # Machine filter checkboxes
    machines_bar = tk.Frame(live_frame, bg='#f0f0f0')
    machines_bar.pack(fill="x", padx=10, pady=(8,0))

    tk.Label(machines_bar, text="Machine Filters:", font=('Segoe UI', 11, 'bold'), bg='#f0f0f0').pack(side=tk.LEFT, padx=(0,10), pady=6)
    live_status_label = tk.Label(live_frame, text="Connecting to live DB...", font=('Segoe UI', 10, 'italic'), bg='#f0f0f0', fg='#666')
    live_status_label.pack(fill='x', padx=10, pady=(0,6))
    global machine_vars
    machine_vars.clear()
    global machine_cbs
    machine_cbs.clear()
    for name in MACHINE_NAMES:
        var = tk.BooleanVar(value=True)
        cb = tk.Checkbutton(
            machines_bar,
            text=name,
            variable=var,
            bg='#f0f0f0',
            font=('Segoe UI', 10, 'bold'),
            padx=10,
            pady=6,
            width=5,
            anchor='w'
        )
        cb.pack(side=tk.LEFT, padx=6, pady=4)
        machine_vars[name] = var
        machine_cbs[name] = cb

    tree_frame = ttk.Frame(live_frame)
    tree_frame.pack(fill="both", expand=True, padx=10, pady=10)

    v_scroll = ttk.Scrollbar(tree_frame, orient="vertical")
    v_scroll.pack(side=tk.RIGHT, fill=tk.Y)
    h_scroll = ttk.Scrollbar(tree_frame, orient="horizontal")
    h_scroll.pack(side=tk.BOTTOM, fill=tk.X)

    live_tree = ttk.Treeview(
        tree_frame,
        yscrollcommand=v_scroll.set,
        xscrollcommand=h_scroll.set,
        height=30
    )
    live_tree.pack(fill="both", expand=True)
    v_scroll.config(command=live_tree.yview)
    h_scroll.config(command=live_tree.xview)

    live_tree.tag_configure("ok", background="#92D050")
    live_tree.tag_configure("low", background="#FFC000")
    live_tree.tag_configure("high", background="#FF6B6B")
    live_tree.tag_configure("flash", background="#87CEEB")

    # Wire checkbox toggles to refresh live view immediately (force rebuild)
    for cb in machine_cbs.values():
        cb.config(command=lambda lt=live_tree: update_live(lt, force_rebuild=True))

    update_live(live_tree)

    # ---------------- DOWNLOAD TAB ----------------
    download_frame = ttk.Frame(notebook)
    notebook.add(download_frame, text="📥 Download Records")

    control_frame = tk.Frame(download_frame, bg='#e8f0fe', relief=tk.RAISED, bd=2)
    control_frame.pack(fill="x", padx=10, pady=10)

    title_label = tk.Label(control_frame, text="📋 Record Export Tool", 
                          font=('Segoe UI', 14, 'bold'), 
                          bg='#e8f0fe', fg='#1a3c6e')
    title_label.pack(pady=5)

    date_frame = tk.Frame(control_frame, bg='#e8f0fe')
    date_frame.pack(pady=10)

    tk.Label(date_frame, text="📅 Start Date:", 
            font=('Segoe UI', 10, 'bold'), 
            bg='#e8f0fe', fg='#1a3c6e').pack(side=tk.LEFT, padx=5)
    
    start_cal = DateEntry(date_frame, width=12, background='darkblue',
                         foreground='white', borderwidth=2,
                         date_pattern='dd-mm-yyyy', font=('Segoe UI', 10))
    start_cal.pack(side=tk.LEFT, padx=5)

    tk.Label(date_frame, text="📅 End Date:", 
            font=('Segoe UI', 10, 'bold'), 
            bg='#e8f0fe', fg='#1a3c6e').pack(side=tk.LEFT, padx=5)
    
    end_cal = DateEntry(date_frame, width=12, background='darkblue',
                       foreground='white', borderwidth=2,
                       date_pattern='dd-mm-yyyy', font=('Segoe UI', 10))
    end_cal.pack(side=tk.LEFT, padx=5)

    # Download machine filter checkboxes
    download_machine_frame = tk.Frame(control_frame, bg='#f0f4ff')
    download_machine_frame.pack(fill="x", padx=10, pady=(0,10))

    tk.Label(download_machine_frame, text="Machine Filter:", 
             font=('Segoe UI', 10, 'bold'), 
             bg='#f0f4ff', fg='#1a3c6e').pack(side=tk.LEFT, padx=5, pady=6)

    global download_machine_vars
    download_machine_vars.clear()
    global download_machine_cbs
    download_machine_cbs.clear()
    for name in MACHINE_NAMES:
        var = tk.BooleanVar(value=True)
        cb = tk.Checkbutton(
            download_machine_frame,
            text=name,
            variable=var,
            bg='#f0f4ff',
            font=('Segoe UI', 10, 'bold'),
            padx=10,
            pady=6,
            width=5,
            anchor='w'
        )
        cb.pack(side=tk.LEFT, padx=6, pady=4)
        download_machine_vars[name] = var
        download_machine_cbs[name] = cb

    button_frame = tk.Frame(control_frame, bg='#e8f0fe')
    button_frame.pack(pady=10)

    def create_styled_button(parent, text, command, color='#4472C4'):
        btn = tk.Button(parent, text=text, command=command,
                       font=('Segoe UI', 10, 'bold'),
                       bg=color, fg='white',
                       padx=20, pady=8,
                       relief=tk.RAISED, bd=2,
                       cursor='hand2')
        return btn

    search_btn = create_styled_button(button_frame, "🔍 Search Records", download_records, '#4472C4')
    search_btn.pack(side=tk.LEFT, padx=5)

    export_button = create_styled_button(button_frame, "📊 Export to Excel", export_previewed_records, '#28A745')
    export_button.pack(side=tk.LEFT, padx=5)
    export_button.config(state="disabled")

    clear_btn = create_styled_button(button_frame, "🗑️ Clear", clear_records, '#DC3545')
    clear_btn.pack(side=tk.LEFT, padx=5)

    status_label = tk.Label(control_frame, text="Ready", 
                           font=('Segoe UI', 10, 'italic'),
                           bg='#e8f0fe', fg='gray')
    status_label.pack(pady=5)

    preview_frame = tk.Frame(download_frame, bg='white', relief=tk.SUNKEN, bd=2)
    preview_frame.pack(fill="both", expand=True, padx=10, pady=10)

    v_scroll2 = ttk.Scrollbar(preview_frame, orient="vertical")
    v_scroll2.pack(side=tk.RIGHT, fill=tk.Y)
    h_scroll2 = ttk.Scrollbar(preview_frame, orient="horizontal")
    h_scroll2.pack(side=tk.BOTTOM, fill=tk.X)

    preview_tree = ttk.Treeview(
        preview_frame,
        yscrollcommand=v_scroll2.set,
        xscrollcommand=h_scroll2.set,
        height=25
    )
    preview_tree.pack(fill="both", expand=True)
    v_scroll2.config(command=preview_tree.yview)
    h_scroll2.config(command=preview_tree.xview)

    preview_tree.tag_configure("ok", background="#92D050")
    preview_tree.tag_configure("low", background="#FFC000")
    preview_tree.tag_configure("high", background="#FF6B6B")

    root.mainloop()

# ---------------- LOGIN ----------------
def handle_login():
    username = username_entry.get()
    password = password_entry.get()

    if username == "admin" and password == "admin":
        login.destroy()
        open_main_window()
    else:
        messagebox.showerror("Login Failed", "Invalid Username or Password")

# ---------------- LOGIN WINDOW (SIMPLIFIED AND FIXED) ----------------
login = tk.Tk()
login.title("Login")
login.geometry("400x350")
login.resizable(False, False)
login.configure(bg='#e8f0fe')

# Create a simple frame
frame = tk.Frame(login, bg='white', bd=2, relief=tk.RAISED)
frame.place(relx=0.5, rely=0.5, anchor='center', width=320, height=280)

# Title at top
title = tk.Label(frame, text="🔐 System Login", font=('Segoe UI', 16, 'bold'), bg='white', fg='#1a3c6e')
title.place(x=160, y=30, anchor='center')

# Username label and entry
user_label = tk.Label(frame, text="Username", font=('Segoe UI', 10), bg='white')
user_label.place(x=160, y=70, anchor='center')

username_entry = tk.Entry(frame, font=('Segoe UI', 11), width=25, relief=tk.SUNKEN, bd=2)
username_entry.place(x=160, y=95, anchor='center')

# Password label and entry
pass_label = tk.Label(frame, text="Password", font=('Segoe UI', 10), bg='white')
pass_label.place(x=160, y=135, anchor='center')

password_entry = tk.Entry(frame, show='*', font=('Segoe UI', 11), width=25, relief=tk.SUNKEN, bd=2)
password_entry.place(x=160, y=160, anchor='center')

# Login Button - PLACED DIRECTLY WITH place()
login_btn = tk.Button(frame, 
                     text="🔓 LOGIN", 
                     command=handle_login,
                     font=('Segoe UI', 12, 'bold'),
                     bg='#4472C4', 
                     fg='white',
                     width=20,
                     relief=tk.RAISED, 
                     bd=2,
                     cursor='hand2')
login_btn.place(x=160, y=205, anchor='center')

# Hover effects
def on_enter(e):
    login_btn.config(bg='#3A5C9E')

def on_leave(e):
    login_btn.config(bg='#4472C4')

login_btn.bind('<Enter>', on_enter)
login_btn.bind('<Leave>', on_leave)

# Hint text
hint = tk.Label(frame, text="Default: admin / admin", font=('Segoe UI', 8, 'italic'), bg='white', fg='#999')
hint.place(x=160, y=245, anchor='center')

# Bind Enter key
def on_enter_key(event):
    handle_login()

login.bind('<Return>', on_enter_key)

login.mainloop()